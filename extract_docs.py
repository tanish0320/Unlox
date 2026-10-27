import docx
import pypdf
import os

files = os.listdir('.')
pdf_file = [f for f in files if f.endswith('.pdf')][0]
prd_file = [f for f in files if 'PRD_FRD' in f and f.endswith('.docx')][0]
wf_file = [f for f in files if 'Workflow' in f and f.endswith('.docx')][0]

print("PDF:", pdf_file)
print("PRD:", prd_file)
print("WF:", wf_file)

def dump_docx(filename, outname):
    doc = docx.Document(filename)
    text = []
    for p in doc.paragraphs:
        if p.text.strip():
            text.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            row_txt = [c.text.strip().replace('\n', ' ') for c in row.cells]
            text.append(' | '.join(row_txt))
    with open(outname, 'w', encoding='utf-8') as f:
        f.write('\n'.join(text))
    print(f'Wrote {outname} ({len(text)} lines)')

def dump_pdf(filename, outname):
    reader = pypdf.PdfReader(filename)
    text = []
    for i, page in enumerate(reader.pages):
        text.append(f'--- PAGE {i+1} ---')
        text.append(page.extract_text() or '')
    with open(outname, 'w', encoding='utf-8') as f:
        f.write('\n'.join(text))
    print(f'Wrote {outname} ({len(reader.pages)} pages)')

dump_pdf(pdf_file, 'problem_statement_text.txt')
dump_docx(prd_file, 'prd_frd_text.txt')
dump_docx(wf_file, 'team_workflow_text.txt')
